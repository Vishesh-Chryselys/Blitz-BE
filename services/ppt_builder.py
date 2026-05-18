import json
import os
import re
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from models.schemas import DocumentReference, PPTRequest
from services.llm import get_generator_llm, get_llm


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

COLORS = {
    "background": RGBColor(250, 248, 241),
    "card": RGBColor(255, 255, 255),
    "ink": RGBColor(43, 35, 25),
    "muted": RGBColor(112, 95, 68),
    "gold": RGBColor(181, 137, 48),
    "gold_soft": RGBColor(244, 230, 187),
    "line": RGBColor(222, 203, 157),
}


# ---------- text helpers ----------


def _safe_text(value: Any, fallback: str = "") -> str:
    if value is None:
        return fallback
    return str(value).strip() or fallback


def _clean_json(content: str) -> str:
    content = content.strip()
    if content.startswith("```json"):
        return content[7:-3].strip()
    if content.startswith("```"):
        return content[3:-3].strip()
    return content


def _parse_json_lenient(raw: str) -> Optional[dict]:
    """Try strict parse first, then recover the largest JSON object in the response."""
    try:
        return json.loads(_clean_json(raw))
    except Exception:
        match = re.search(r"\{[\s\S]*\}", raw or "")
        if match:
            try:
                return json.loads(match.group(0))
            except Exception:
                return None
        return None


def _shorten(value: str, max_chars: int) -> str:
    value = _safe_text(value)
    return value if len(value) <= max_chars else value[: max_chars - 1].rstrip() + "..."


def _one_line(value: str, max_chars: int) -> str:
    return _shorten(re.sub(r"\s+", " ", _safe_text(value)), max_chars)


def _compact_bullets(values: List[Any], max_items: int = 4, max_chars: int = 95) -> List[str]:
    bullets: List[str] = []
    for value in values or []:
        text = _one_line(value, max_chars)
        if text:
            bullets.append(text)
        if len(bullets) >= max_items:
            break
    return bullets


def _chunks(values: List[Any], size: int) -> List[List[Any]]:
    return [values[i:i + size] for i in range(0, len(values), size)]


def _reference_payload(references: List[DocumentReference]) -> List[Dict[str, Any]]:
    payload = []
    for idx, ref in enumerate(references, start=1):
        meta = ref.metadata or {}
        payload.append({
            "doc_id": f"Doc {idx}",
            "file_name": _one_line(ref.file_name, 90),
            "summary": _one_line(ref.summary, 240),
            "client": meta.get("client_name") or meta.get("client_project") or "N/A",
            "poc": meta.get("pocs") or meta.get("created_by") or "Unknown",
            "objective": meta.get("business_objective") or "",
            "approach": meta.get("approach") or "",
            "datasets": meta.get("datasets_used") or "",
            "outcome": meta.get("key_outcome") or "",
            "topic": meta.get("topic") or "",
            "brand": meta.get("brand") or "",
        })
    return payload


# ---------- planning: intent classifier ----------


INTENTS = {"capability", "client_showcase", "proposal_scaffold"}


def _classify_intent(request: PPTRequest, refs: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Cheap Haiku call that picks the deck shape and emits title/subtitle/hero."""
    llm = get_llm()
    ref_clients = sorted({r["client"] for r in refs if r["client"] and r["client"] != "N/A"})
    prompt = f"""You are BLITZ Deck Architect, classifying the kind of sales deck to produce.

Choose ONE intent:
- "capability"         : a generic Chryselys capability / methodology overview (no specific named engagement).
- "client_showcase"    : showcases the work Chryselys did for ONE specific named pharma client.
- "proposal_scaffold"  : a forward-looking proposal / pitch for a prospect or new engagement.

Return ONLY valid JSON:
{{
  "intent": "capability|client_showcase|proposal_scaffold",
  "deck_title": "punchy <8-word title",
  "subtitle": "<14-word strapline that frames the deck>",
  "primary_client": "<client name if intent is client_showcase, else null>",
  "hero_outcome": "<one quotable result or stat pulled from the references, else empty string>"
}}

Topic: {request.topic}
User question: {request.user_query or 'N/A'}
Reference clients on file: {', '.join(ref_clients) or 'None'}

RAG answer (first 1500 chars):
{(request.content or '')[:1500]}
"""
    try:
        raw = llm.invoke(prompt).content
        parsed = _parse_json_lenient(raw)
        if parsed and parsed.get("intent") in INTENTS:
            return {
                "intent": parsed["intent"],
                "deck_title": _one_line(parsed.get("deck_title") or request.topic, 70),
                "subtitle": _one_line(parsed.get("subtitle") or "Generated by BLITZ Enterprise Intelligence", 120),
                "primary_client": parsed.get("primary_client") or None,
                "hero_outcome": _one_line(parsed.get("hero_outcome") or "", 110),
            }
    except Exception as e:
        print(f"Intent classification failed, defaulting to capability: {e}")

    return {
        "intent": "capability",
        "deck_title": _one_line(request.topic or "BLITZ Capability Deck", 70),
        "subtitle": "Generated by BLITZ Enterprise Intelligence",
        "primary_client": None,
        "hero_outcome": _one_line(refs[0]["outcome"], 110) if refs else "",
    }


# ---------- planning: per-section content ----------


SECTION_BRIEFS: Dict[str, Dict[str, tuple]] = {
    "exec_summary": {
        "capability":        ("Capability Overview", "Decision Brief",
            "Summarize Chryselys' capability in this area. 3 board-level bullets, each <=16 words. Lead with what we do, then who we serve, then differentiators."),
        "client_showcase":   ("Engagement at a Glance", "Decision Brief",
            "Summarize the engagement for the named client. 3 bullets: what we were asked to do, what we delivered, and the outcome."),
        "proposal_scaffold": ("Business Context", "Proposal Brief",
            "Frame the prospect's business problem and why Chryselys can solve it. 3 bullets, executive tone."),
    },
    "approach": {
        "capability":        ("Our Methodology", "How we deliver",
            "5 short steps describing our methodology in this capability area. Each step <=14 words, action-led."),
        "client_showcase":   ("How We Approached This Work", "Execution storyline",
            "5 steps describing the actual delivery sequence for this client. Each step <=14 words."),
        "proposal_scaffold": ("Proposed Approach", "Path to value",
            "5 steps describing what we would do for the prospect. Each step <=14 words, outcome-linked."),
    },
    "next_steps": {
        "capability":        ("Engagement Model", "How to start",
            "3 pragmatic ways a prospect can engage with this capability (pilot, workshop, embedded analytics, etc)."),
        "client_showcase":   ("How to Leverage This", "Reusable patterns",
            "3 ways another pharma client could benefit from the same playbook."),
        "proposal_scaffold": ("Recommended Next Steps", "Path to kickoff",
            "3 concrete actions to move the proposal forward (discovery, scoping, data access)."),
    },
}


def _section_prompt(kind: str, intent: str, request: PPTRequest, intent_meta: Dict[str, Any], refs: List[Dict[str, Any]]) -> str:
    heading, kicker, directive = SECTION_BRIEFS[kind][intent]
    refs_blob = json.dumps(refs[:5], indent=2)[:2800]
    return f"""You are BLITZ Deck Architect writing one slide of a Chryselys consulting deck.

Deck context:
  Intent:         {intent}
  Deck title:     {intent_meta['deck_title']}
  Primary client: {intent_meta.get('primary_client') or 'N/A'}
  Topic:          {request.topic}
  User question:  {request.user_query or 'N/A'}

Slide you are writing:
  Heading:   {heading}
  Kicker:    {kicker}
  Directive: {directive}

Rules:
- Use ONLY facts grounded in the RAG answer and references. Never invent client names, POCs, or numbers.
- No bullet exceeds 18 words. Be specific and concrete - no fluff.
- Speaker notes: 2-3 sentences a sales rep would actually say out loud while showing the slide.

Return ONLY valid JSON:
{{
  "bullets": ["...", "...", "..."],
  "speaker_notes": "..."
}}

RAG answer (first 3500 chars):
{(request.content or '')[:3500]}

References:
{refs_blob}
"""


def _ref_card_prompt(ref: Dict[str, Any], intent: str, request: PPTRequest, intent_meta: Dict[str, Any]) -> str:
    return f"""You are BLITZ Deck Architect writing ONE Chryselys engagement card for a single slide.

Deck context:
  Intent:     {intent}
  Deck title: {intent_meta['deck_title']}
  Topic:      {request.topic}

Reference (one indexed document):
{json.dumps(ref, indent=2)[:2200]}

RAG answer (first 2500 chars):
{(request.content or '')[:2500]}

Rules:
- Stay grounded in the reference fields. If a field is empty, write "Not specified in the indexed material" - never invent.
- heading <= 9 words, scoped to this single engagement.
- objective / approach / datasets / outcome each <= 22 words.
- speaker_notes: 2-3 conversational sentences a sales rep would say while presenting this slide.

Return ONLY valid JSON:
{{
  "heading":  "<scoped slide title>",
  "objective":"<the business problem we tackled>",
  "approach": "<what we actually did>",
  "datasets": "<data / tools used>",
  "outcome":  "<measurable result or deliverable>",
  "speaker_notes": "..."
}}
"""


def _invoke_section_llm(prompt: str) -> Dict[str, Any]:
    try:
        raw = get_generator_llm().invoke(prompt).content
        return _parse_json_lenient(raw) or {}
    except Exception as e:
        print(f"Section LLM call failed: {e}")
        return {}


def _generate_sections(request: PPTRequest, intent_meta: Dict[str, Any], refs: List[Dict[str, Any]]) -> Dict[str, Any]:
    intent = intent_meta["intent"]
    section_kinds = ["exec_summary", "approach", "next_steps"]
    top_refs = refs[:3]

    jobs: Dict[str, str] = {kind: _section_prompt(kind, intent, request, intent_meta, refs) for kind in section_kinds}
    for i, ref in enumerate(top_refs):
        jobs[f"ref_{i}"] = _ref_card_prompt(ref, intent, request, intent_meta)

    results: Dict[str, Dict[str, Any]] = {}
    if jobs:
        with ThreadPoolExecutor(max_workers=min(6, len(jobs))) as ex:
            future_map = {ex.submit(_invoke_section_llm, prompt): name for name, prompt in jobs.items()}
            for fut in as_completed(future_map):
                results[future_map[fut]] = fut.result()

    sections: Dict[str, Any] = {}
    for kind in section_kinds:
        heading, kicker, _ = SECTION_BRIEFS[kind][intent]
        payload = results.get(kind) or {}
        sections[kind] = {
            "heading": heading,
            "kicker": kicker,
            "bullets": _compact_bullets(payload.get("bullets", []), max_items=5, max_chars=140),
            "speaker_notes": _one_line(payload.get("speaker_notes", ""), 600),
        }

    ref_cards: List[Dict[str, Any]] = []
    for i, ref in enumerate(top_refs):
        payload = results.get(f"ref_{i}") or {}
        ref_cards.append({
            "doc_id": ref["doc_id"],
            "client": ref["client"],
            "poc": ref["poc"],
            "heading": _one_line(payload.get("heading") or ref["topic"] or ref["file_name"], 70),
            "objective": _one_line(payload.get("objective") or ref["objective"] or ref["summary"], 220),
            "approach": _one_line(payload.get("approach") or ref["approach"], 220),
            "datasets": _one_line(payload.get("datasets") or ref["datasets"], 200),
            "outcome": _one_line(payload.get("outcome") or ref["outcome"], 220),
            "speaker_notes": _one_line(payload.get("speaker_notes", ""), 600),
        })
    sections["ref_cards"] = ref_cards

    sections["cover"] = {
        "speaker_notes": _one_line(
            f"Open by framing this as Chryselys' synthesized intelligence on {request.topic}. "
            f"Hero: {intent_meta.get('hero_outcome') or 'evidence-backed Chryselys capability with cited proof points'}.",
            500,
        ),
    }
    sections["appendix"] = {
        "speaker_notes": "Walk through every source document this deck drew on. Each row is a real Chryselys deliverable you can follow up on with the listed POC.",
    }

    return sections


# ---------- pptx primitives ----------


def _set_background(slide, color=COLORS["background"]):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _add_text(slide, text, x, y, w, h, size=18, color=None, bold=False, align=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = Inches(0.01)
    p = frame.paragraphs[0]
    p.text = _safe_text(text)
    if align:
        p.alignment = align
    run = p.runs[0] if p.runs else p.add_run()
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Aptos"
    run.font.color.rgb = color or COLORS["ink"]
    return box


def _add_bullets(slide, bullets, x, y, w, h, size=15, color=None, max_items=5, max_chars=140):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.word_wrap = True
    for idx, bullet in enumerate(_compact_bullets(bullets, max_items=max_items, max_chars=max_chars)):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = f"- {bullet}"
        p.level = 0
        p.space_after = Pt(5)
        p.font.size = Pt(size)
        p.font.name = "Aptos"
        p.font.color.rgb = color or COLORS["ink"]
    return box


def _add_header(slide, title: str, kicker: str):
    _add_text(slide, kicker.upper(), Inches(0.65), Inches(0.34), Inches(6.5), Inches(0.25), 8, COLORS["gold"], True)
    _add_text(slide, title, Inches(0.65), Inches(0.62), Inches(11.5), Inches(0.62), 24, COLORS["ink"], True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.32), Inches(12.0), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["line"]
    line.line.fill.background()


def _add_footer(slide, page: int):
    _add_text(slide, "BLITZ Enterprise Intelligence | Generated from cited RAG synthesis", Inches(0.65), Inches(7.08), Inches(7.0), Inches(0.2), 7, COLORS["muted"])
    _add_text(slide, str(page).zfill(2), Inches(12.25), Inches(7.05), Inches(0.45), Inches(0.22), 8, COLORS["gold"], True, PP_ALIGN.RIGHT)


def _set_notes(slide, text: str):
    if not text:
        return
    notes_frame = slide.notes_slide.notes_text_frame
    notes_frame.text = _safe_text(text)


def _new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ---------- slide renderers ----------


def _slide_cover(prs, plan):
    slide = _new_slide(prs)
    _set_background(slide)
    _add_text(slide, "BLITZ", Inches(0.72), Inches(0.45), Inches(1.7), Inches(0.3), 14, COLORS["gold"], True)
    _add_text(slide, _one_line(plan["deck_title"], 58), Inches(0.72), Inches(1.75), Inches(7.7), Inches(1.3), 31, COLORS["ink"], True)
    _add_text(slide, _one_line(plan["subtitle"], 100), Inches(0.78), Inches(3.18), Inches(7.4), Inches(0.55), 14, COLORS["muted"])

    hero = plan.get("hero_outcome")
    if hero:
        band = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(4.05), Inches(7.7), Inches(1.15))
        band.fill.solid()
        band.fill.fore_color.rgb = COLORS["gold_soft"]
        band.line.color.rgb = COLORS["line"]
        _add_text(slide, "HERO OUTCOME", Inches(0.95), Inches(4.18), Inches(3.5), Inches(0.22), 8, COLORS["gold"], True)
        _add_text(slide, hero, Inches(0.95), Inches(4.42), Inches(7.3), Inches(0.65), 14, COLORS["ink"], True)
    else:
        bullets = plan.get("sections", {}).get("exec_summary", {}).get("bullets", [])
        _add_bullets(slide, bullets, Inches(0.85), Inches(4.08), Inches(7.4), Inches(1.6), 12, max_items=3, max_chars=120)

    orb = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.95), Inches(1.85), Inches(3.65), Inches(3.65))
    orb.fill.solid()
    orb.fill.fore_color.rgb = COLORS["gold_soft"]
    orb.line.color.rgb = COLORS["line"]
    _add_text(slide, plan["intent"].replace("_", " ").upper(), Inches(9.05), Inches(2.95), Inches(3.45), Inches(0.55), 16, COLORS["gold"], True, PP_ALIGN.CENTER)
    _add_text(slide, "Cited synthesis · Source mapping · Proposal-ready", Inches(9.05), Inches(3.65), Inches(3.45), Inches(0.55), 10, COLORS["muted"], False, PP_ALIGN.CENTER)
    _add_footer(slide, 1)
    _set_notes(slide, plan["sections"]["cover"].get("speaker_notes", ""))


def _slide_bullets_with_card(prs, plan, section_key: str, page: int, card_title: str, card_body: str):
    section = plan["sections"][section_key]
    slide = _new_slide(prs)
    _set_background(slide)
    _add_header(slide, section["heading"], section["kicker"])
    _add_bullets(slide, section["bullets"], Inches(0.9), Inches(1.75), Inches(7.4), Inches(4.5), 15, max_items=6, max_chars=140)

    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.55), Inches(1.75), Inches(4.25), Inches(4.45))
    card.fill.solid()
    card.fill.fore_color.rgb = COLORS["card"]
    card.line.color.rgb = COLORS["line"]
    _add_text(slide, card_title, Inches(8.78), Inches(1.95), Inches(3.85), Inches(0.35), 11, COLORS["gold"], True)
    _add_text(slide, card_body, Inches(8.78), Inches(2.4), Inches(3.85), Inches(3.55), 10, COLORS["muted"])
    _add_footer(slide, page)
    _set_notes(slide, section.get("speaker_notes", ""))


def _slide_steps(prs, plan, section_key: str, page: int):
    section = plan["sections"][section_key]
    slide = _new_slide(prs)
    _set_background(slide)
    _add_header(slide, section["heading"], section["kicker"])
    steps = _compact_bullets(section["bullets"], max_items=5, max_chars=80)
    left = Inches(0.95)
    top = Inches(2.2)
    gap = Inches(2.35)
    for idx, step in enumerate(steps):
        x = left + Inches(idx * 2.35)
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, top, Inches(0.58), Inches(0.58))
        circle.fill.solid()
        circle.fill.fore_color.rgb = COLORS["gold"]
        circle.line.fill.background()
        _add_text(slide, str(idx + 1), x + Inches(0.16), top + Inches(0.09), Inches(0.25), Inches(0.24), 12, RGBColor(255, 255, 255), True, PP_ALIGN.CENTER)
        if idx < len(steps) - 1:
            connector = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(0.68), top + Inches(0.28), gap - Inches(0.15), Inches(0.025))
            connector.fill.solid()
            connector.fill.fore_color.rgb = COLORS["line"]
            connector.line.fill.background()
        _add_text(slide, step, x - Inches(0.18), top + Inches(0.86), Inches(2.05), Inches(1.6), 11, COLORS["ink"], True, PP_ALIGN.CENTER)
    _add_footer(slide, page)
    _set_notes(slide, section.get("speaker_notes", ""))


def _slide_ref_card(prs, ref: Dict[str, Any], page: int):
    slide = _new_slide(prs)
    _set_background(slide)
    kicker = f"Engagement | {_one_line(ref.get('client', 'N/A'), 30)}"
    _add_header(slide, _one_line(ref.get("heading", ref.get("doc_id", "Engagement")), 65), kicker)

    strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.45), Inches(12.0), Inches(0.55))
    strip.fill.solid()
    strip.fill.fore_color.rgb = COLORS["gold_soft"]
    strip.line.color.rgb = COLORS["line"]
    _add_text(slide, f"POC: {_one_line(ref.get('poc', 'Unknown'), 40)}", Inches(0.85), Inches(1.55), Inches(8.0), Inches(0.35), 11, COLORS["ink"], True)
    _add_text(slide, _one_line(ref.get("doc_id", ""), 16), Inches(11.3), Inches(1.55), Inches(1.25), Inches(0.35), 10, COLORS["gold"], True, PP_ALIGN.RIGHT)

    facets = [
        ("Business Objective", ref.get("objective")),
        ("Approach", ref.get("approach")),
        ("Datasets & Tools", ref.get("datasets")),
        ("Outcome", ref.get("outcome")),
    ]
    grid_positions = [(0.65, 2.25), (6.95, 2.25), (0.65, 4.6), (6.95, 4.6)]
    for (title, body), (x, y) in zip(facets, grid_positions):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(6.0), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = COLORS["card"]
        card.line.color.rgb = COLORS["line"]
        _add_text(slide, title.upper(), Inches(x + 0.25), Inches(y + 0.15), Inches(5.5), Inches(0.3), 9, COLORS["gold"], True)
        _add_text(slide, _one_line(body or "Not specified in the indexed material", 260), Inches(x + 0.25), Inches(y + 0.5), Inches(5.5), Inches(1.6), 11, COLORS["muted"])

    _add_footer(slide, page)
    _set_notes(slide, ref.get("speaker_notes", ""))


def _slide_appendix(prs, evidence: List[Dict[str, Any]], page: int, part: int, total_parts: int, speaker_notes: str):
    slide = _new_slide(prs)
    _set_background(slide)
    title = "Appendix: Referenced Documents" if total_parts == 1 else f"Appendix: Referenced Documents ({part}/{total_parts})"
    _add_header(slide, title, "Full source list")
    bullets = [
        f"{item.get('doc_id', f'Doc {idx + 1}')}: {_one_line(item.get('file_name', 'Source'), 48)} | Client: {_one_line(item.get('client', 'N/A'), 20)} | POC: {_one_line(item.get('poc', 'Unknown'), 20)}"
        for idx, item in enumerate(evidence)
    ] or ["No source references were provided with this deck request."]
    _add_bullets(slide, bullets, Inches(0.9), Inches(1.65), Inches(11.5), Inches(4.8), 12, max_items=14, max_chars=170)
    _add_footer(slide, page)
    _set_notes(slide, speaker_notes)


# ---------- shape definitions ----------


SHAPES: Dict[str, List[tuple]] = {
    "capability": [
        ("cover",),
        ("exec_summary", "What this means",
            "Use this deck as the bridge between the chat answer and the cited Chryselys deliverables - both capability framing and proof points."),
        ("approach",),
        ("ref_card", 0),
        ("ref_card", 1),
        ("ref_card", 2),
        ("next_steps", "BLITZ output advantage",
            "Leave-behind for capability conversations. Every bullet ties back to a cited Chryselys document."),
        ("appendix",),
    ],
    "client_showcase": [
        ("cover",),
        ("exec_summary", "Why this engagement matters",
            "Anchors the work Chryselys did for the named client and shows how that experience transfers to similar pharma asks."),
        ("ref_card", 0),
        ("ref_card", 1),
        ("ref_card", 2),
        ("approach",),
        ("next_steps", "How to reuse",
            "Patterns from this engagement that can be lifted directly into the next pharma client conversation."),
        ("appendix",),
    ],
    "proposal_scaffold": [
        ("cover",),
        ("exec_summary", "Why us, why now",
            "Frames the prospect's problem and positions Chryselys as the best-equipped partner to solve it, backed by cited prior work."),
        ("approach",),
        ("ref_card", 0),
        ("ref_card", 1),
        ("ref_card", 2),
        ("next_steps", "From proposal to kickoff",
            "Concrete steps to convert this proposal into a signed engagement."),
        ("appendix",),
    ],
}


# ---------- top-level entry ----------


def build_ppt(request: PPTRequest) -> str:
    references = _reference_payload(request.references)
    intent_meta = _classify_intent(request, references)
    sections = _generate_sections(request, intent_meta, references)
    plan = {**intent_meta, "sections": sections}

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    shape = SHAPES.get(intent_meta["intent"], SHAPES["capability"])
    ref_cards: List[Dict[str, Any]] = sections.get("ref_cards", [])
    evidence_pages = _chunks(references[:14], 7) or [[]]

    page = 1
    for entry in shape:
        kind = entry[0]
        extras = entry[1:]

        if kind == "cover":
            _slide_cover(prs, plan)
            page = 2
        elif kind in ("exec_summary", "next_steps"):
            card_title, card_body = extras
            _slide_bullets_with_card(prs, plan, kind, page, card_title, card_body)
            page += 1
        elif kind == "approach":
            _slide_steps(prs, plan, kind, page)
            page += 1
        elif kind == "ref_card":
            idx = extras[0]
            if idx < len(ref_cards):
                _slide_ref_card(prs, ref_cards[idx], page)
                page += 1
        elif kind == "appendix":
            speaker_notes = sections["appendix"].get("speaker_notes", "")
            for part, ev_chunk in enumerate(evidence_pages, start=1):
                _slide_appendix(prs, ev_chunk, page, part, len(evidence_pages), speaker_notes)
                page += 1

    os.makedirs("exports", exist_ok=True)
    safe_topic = re.sub(r"[^a-zA-Z0-9_-]+", "_", request.topic).strip("_")[:36] or "blitz_deck"
    file_path = os.path.join("exports", f"{safe_topic}_{uuid.uuid4().hex[:8]}.pptx")
    prs.save(file_path)
    return file_path
