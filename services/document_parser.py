import os
import fitz  # PyMuPDF
from pptx import Presentation
import pandas as pd
import docx
from typing import Dict, Any

def parse_pdf(file_path: str) -> Dict[str, Any]:
    text = ""
    with fitz.open(file_path) as doc:
        for page in doc:
            text += page.get_text() + "\n"
    return {"text": text, "type": "pdf"}

def parse_pptx(file_path: str) -> Dict[str, Any]:
    prs = Presentation(file_path)
    slides_data = []
    full_text = ""
    
    for i, slide in enumerate(prs.slides):
        slide_text = ""
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                slide_text += shape.text + "\n"
        slides_data.append({"slide_number": i + 1, "content": slide_text.strip()})
        full_text += slide_text + "\n"
        
    return {"text": full_text, "slides": slides_data, "type": "pptx"}

MAX_ROWS = 500  # Cap to avoid token explosion on massive spreadsheets

def parse_excel(file_path: str) -> Dict[str, Any]:
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".xlsb":
            df = pd.read_excel(file_path, engine="pyxlsb", nrows=MAX_ROWS)
        else:
            df = pd.read_excel(file_path, nrows=MAX_ROWS)
    except Exception as e:
        return {"text": "", "error": str(e), "type": ext}
    
    total_hint = f"[Note: Showing first {MAX_ROWS} rows of spreadsheet '{os.path.basename(file_path)}']\n"
    text = total_hint + df.to_csv(index=False)
    return {"text": text, "type": ext.lstrip(".")}

def parse_docx(file_path: str) -> Dict[str, Any]:
    doc = docx.Document(file_path)
    text = "\n".join([para.text for para in doc.paragraphs])
    return {"text": text, "type": "docx"}

def parse_csv(file_path: str) -> Dict[str, Any]:
    df = pd.read_csv(file_path, nrows=MAX_ROWS)
    total_hint = f"[Note: Showing first {MAX_ROWS} rows of CSV '{os.path.basename(file_path)}']\n"
    text = total_hint + df.to_csv(index=False)
    return {"text": text, "type": "csv"}

def parse_document(file_path: str) -> Dict[str, Any]:
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == ".pdf":
            return parse_pdf(file_path)
        elif ext == ".pptx":
            return parse_pptx(file_path)
        elif ext in [".xlsx", ".xls", ".xlsb"]:
            return parse_excel(file_path)
        elif ext in [".docx", ".doc"]:
            return parse_docx(file_path)
        elif ext == ".csv":
            return parse_csv(file_path)
        elif ext in [".txt", ".md"]:
            with open(file_path, "r", encoding="utf-8") as f:
                return {"text": f.read(), "type": "txt"}
        else:
            return {"text": f"Unsupported file format: {ext}", "type": "unknown"}
    except Exception as e:
        print(f"Error parsing {file_path}: {e}")
        return {"text": "", "error": str(e), "type": ext}
