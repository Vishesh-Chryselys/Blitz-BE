import os
from fastapi import APIRouter
from fastapi.responses import FileResponse
from pydantic import BaseModel
from pptx import Presentation
from services.llm import get_llm

router = APIRouter()

class PPTRequest(BaseModel):
    topic: str
    content: str # The synthesized AI answer to convert to slides

@router.post("/")
async def generate_ppt(request: PPTRequest):
    """AGENT 3 (PPT Designer): Generates a structured PowerPoint using structured JSON LLM output."""
    llm = get_llm()
    
    # Prompt the LLM to act as a Structured PPT Design Agent
    prompt = f"""You are the Advanced PPT Design Agent. 
Analyze the provided business analysis content and convert it into a structured, professional slide deck on the topic '{request.topic}'.

Structure Requirements:
1. Generate 3 to 4 content slides capturing the core details, strategic use cases, and insights.
2. The FINAL slide MUST strictly be an 'Appendix' listing all referenced document names/sources found in the content.

Format your response STRICTLY as a valid JSON array of objects. Do not include markdown code block syntax (like ```json).
Each slide object MUST have:
- "title": A short, high-impact slide title (string)
- "bullets": An array of 3 to 4 concise, high-value bullet points (strings)

Content to transform:
{request.content}
"""
    
    try:
        res = llm.invoke(prompt).content.strip()
        
        # Clean potential markdown wrapping
        if res.startswith("```json"):
            res = res[7:-3].strip()
        elif res.startswith("```"):
            res = res[3:-3].strip()
            
        import json
        slides_data = json.loads(res)
    except Exception as e:
        print(f"JSON Slide Parsing failed: {e}. Falling back to default layout.")
        slides_data = [
            {"title": "Overview", "bullets": ["AI-synthesized business strategy.", "Detailed metric evaluation."]},
            {"title": "Appendix: Referenced Assets", "bullets": ["Internal corporate knowledge store."]}
        ]
    
    prs = Presentation()
    
    # 1. Elegant Title Slide
    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = request.topic
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = "Chryselys Sales Enablement Intelligence Portal\nGenerated dynamically by BLITZ AI Agent"
    
    # 2. Dynamic Content Slides & Appendix
    bullet_layout = prs.slide_layouts[1]
    
    for slide_info in slides_data:
        title_text = slide_info.get("title", "Insight Slide")
        bullets = slide_info.get("bullets", [])
        
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = title_text
        
        if len(slide.placeholders) > 1 and bullets:
            slide.placeholders[1].text = "\n".join([f"• {b}" for b in bullets])
            
    os.makedirs("exports", exist_ok=True)
    file_path = f"exports/{request.topic.replace(' ', '_')[:20]}.pptx"
    prs.save(file_path)
    
    return FileResponse(
        path=file_path, 
        filename=f"{request.topic[:20]}.pptx", 
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
